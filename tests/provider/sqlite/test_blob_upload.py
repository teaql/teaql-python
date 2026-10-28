"""Real SQLite/File upload confirmation; parser remains test-only."""

from datetime import datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from shutil import which
from threading import Event
from urllib.request import Request, urlopen
from uuid import uuid4
import asyncio
import os
import time

import aiosqlite
import pytest

from teaql.core.blob_ref import BlobFieldSpec
from teaql.core.entity import EntityKey
from teaql.core.meta import EntityDescriptor, PropertyDescriptor
from teaql.core.mutation import InsertCommand, UpdateCommand
from teaql.core.object_storage import ObjectContent, ObjectKey, PutOptions, WriteMode
from teaql.core.value import DataType
from teaql.provider.sqlite import SimpleSchemaProvider
from teaql.provider.sqlite.dialect import SqliteDialect
from teaql.provider.sqlite.transport import SqliteTransport
from teaql.runtime import RuntimeModule
from teaql.runtime.view_preparation import (
    AuthorizedBlobSlot, AuthorizedResponseGraph, ViewOptions,
    ViewPreparationError, ViewPreparationService,
)
from teaql.sql.executor import SqlDataServiceExecutor
from teaql.sql.types import CompiledQuery
from teaql_blob_reference import (
    BlobBindingChange, BlobBindingTarget, BlobRetentionPolicy, BlobUploadAuthorizationState,
    BlobUploadInput, BlobUploadTarget,
    BlobVerifiedContent, FileObjectStorage, SqliteBlobCatalog, SqliteBlobResponseContributor,
    SqliteBlobMaintenanceService, SqliteBlobUploadService,
)


class UploadAuthority:
    def __init__(self, extensions=("pdf",)):
        self.current_scope = "tenant-a"
        self.allowed = True
        self.extensions = extensions

    def scope(self, _context):
        return self.current_scope

    def field(self, _context, _target):
        return BlobFieldSpec("Asset", "document", self.extensions, True, "default")

    async def authorize(self, _context, target, action, _input):
        return (self.allowed and target.entity == "Asset" and target.field == "document"
                and action in ("Begin", "Write", "Confirm"))

    async def current_policy_on_lease(self, _context, lease, _target):
        assert lease is not None
        return BlobRetentionPolicy("default", "server_stream", 86400, 604800, 3600)


class BindingAuthority:
    def scope(self, _context):
        return "tenant-a"

    async def authorize(self, _context, target, _previous, _next, action):
        return target.entity == "Asset" and target.field == "document" and action == "Create"

    async def current_policy_on_lease(self, _context, lease, _target):
        assert lease is not None
        return BlobRetentionPolicy("default", "server_stream", 86400, 604800, 3600)


class FixturePDFVerifier:
    def __init__(self, early=False):
        self.early = early
        self.calls = 0

    def verify(self, stream, extension, maximum_bytes):
        self.calls += 1
        if self.early:
            return BlobVerifiedContent(1, "0" * 64, "application/pdf")
        data = stream.read(maximum_bytes + 1)
        assert len(data) <= maximum_bytes and extension == "pdf"
        assert data.startswith(b"%PDF-") and b"%%EOF" in data
        assert stream.read(1) == b""
        return BlobVerifiedContent(len(data), sha256(data).hexdigest(), "application/pdf")


@pytest.mark.asyncio
async def test_upload_begin_commit_unknown_replays_original_live_session(tmp_path):
    path = str(tmp_path / "begin-reply-lost.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    actor = RuntimeModule.new().into_context()
    authority = UploadAuthority()
    now = [10000]
    service = SqliteBlobUploadService(catalog, authority, FixturePDFVerifier(), 1000,
                                      now_ms=lambda: now[0])
    target = BlobUploadTarget("Asset", "1", "document")
    upload = BlobUploadInput("report.pdf", 25)
    begin_sql = catalog.transport.begin_sql

    class CommitThenFail:
        def __init__(self, tx):
            self.tx = tx

        def __getattr__(self, name):
            return getattr(self.tx, name)

        async def commit_sql(self):
            await self.tx.commit_sql()
            raise OSError("PRIVATE-BEGIN-REPLY-CANARY")

    async def uncertain_begin():
        return CommitThenFail(await begin_sql())

    catalog.transport.begin_sql = uncertain_begin
    try:
        with pytest.raises(RuntimeError, match="^BLOB_CATALOG_FAILED$") as failure:
            await service.begin(actor, target, upload, "lost-begin")
        assert "CANARY" not in str(failure.value)
    finally:
        catalog.transport.begin_sql = begin_sql

    async with aiosqlite.connect(path) as db:
        async with db.execute("""SELECT u.upload_id,u.key,u.expires_ms,c.status,c.upload_mode
            FROM teaql_blob_upload u JOIN teaql_blob_catalog c ON c.key=u.key
            WHERE u.scope=? AND u.replay_key=?""", ("tenant-a", "lost-begin")) as cursor:
            rows = await cursor.fetchall()
    assert len(rows) == 1
    upload_id, key, expires_ms, status, mode = rows[0]
    assert status == "initiated" and mode == "server_stream"

    authority.allowed = False
    with pytest.raises(RuntimeError, match="^BLOB_UPLOAD_DENIED$"):
        await service.begin(actor, target, upload, "lost-begin")
    authority.allowed = True
    with pytest.raises(RuntimeError, match="^BLOB_IDEMPOTENCY_CONFLICT$"):
        await service.begin(actor, target, BlobUploadInput("changed.pdf", 25), "lost-begin")
    replay = await service.begin(actor, target, upload, "lost-begin")
    assert (replay.upload_id, replay.reference.key,
            int(replay.expires_at.timestamp() * 1000)) == (upload_id, key, expires_ms)
    now[0] = expires_ms
    with pytest.raises(RuntimeError, match="^BLOB_UPLOAD_EXPIRED$"):
        await service.begin(actor, target, upload, "lost-begin")
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT COUNT(*) FROM teaql_blob_upload") as cursor:
            assert (await cursor.fetchone())[0] == 1


@pytest.mark.asyncio
async def test_upload_write_closes_source_once_before_promotion_and_replay(tmp_path):
    class CountedSource(BytesIO):
        def __init__(self, data, *, fail_close=False):
            super().__init__(data)
            self.close_count = 0
            self.fail_close = fail_close

        def close(self):
            self.close_count += 1
            super().close()
            if self.fail_close:
                raise OSError("PRIVATE-UPLOAD-SOURCE-CLOSE-CANARY")
            if self.close_count > 1:
                raise OSError("PRIVATE-DOUBLE-CLOSE-CANARY")

    path = str(tmp_path / "write-source-close.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    actor = RuntimeModule.new().into_context().with_object_storage(
        FileObjectStorage(tmp_path / "private", os.urandom(32)))
    actor.with_object_storage_observer(lambda event: None)
    service = SqliteBlobUploadService(catalog, UploadAuthority(), FixturePDFVerifier(), 10_000)
    data = b"%PDF-1.7\nsource close before promotion\n%%EOF"
    session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                  BlobUploadInput("report.pdf", len(data)), "source-close")
    failed = CountedSource(data, fail_close=True)
    with pytest.raises(RuntimeError, match="^BLOB_STORE_FAILED$"):
        await service.write(actor, session.upload_id, failed)
    assert failed.close_count == 1
    async with aiosqlite.connect(path) as db:
        async with db.execute("""SELECT u.state,c.status FROM teaql_blob_upload u
            JOIN teaql_blob_catalog c ON c.key=u.key WHERE u.upload_id=?""",
                              (session.upload_id,)) as cursor:
            assert await cursor.fetchone() == ("initiated", "initiated")
    retry = CountedSource(data)
    assert await service.write(actor, session.upload_id, retry) == len(data)
    assert retry.close_count == 1
    assert await service.complete(actor, session.upload_id) == session.reference
    replay = CountedSource(data)
    assert await service.write(actor, session.upload_id, replay) == len(data)
    assert replay.close_count == 1


@pytest.mark.asyncio
async def test_upload_confirm_close_failure_rolls_back_and_retries(tmp_path):
    class CloseFailStream:
        def __init__(self, source):
            self.source = source

        def read(self, size=-1):
            return self.source.read(size)

        def close(self):
            self.source.close()
            raise OSError("PRIVATE-OBJECT-CLOSE-CANARY")

    class CloseFailStore(FileObjectStorage):
        fail_close = True

        def get(self, invocation, key):
            content = super().get(invocation, key)
            if self.fail_close:
                return ObjectContent(content.metadata, CloseFailStream(content.stream))
            return content

    path = str(tmp_path / "confirm-close-failure.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    store = CloseFailStore(tmp_path / "private", os.urandom(32))
    actor = RuntimeModule.new().into_context().with_object_storage(store)
    actor.with_object_storage_observer(lambda event: None)
    service = SqliteBlobUploadService(catalog, UploadAuthority(), FixturePDFVerifier(), 10_000)
    target = BlobUploadTarget("Asset", "1", "document")
    data = b"%PDF-1.7\nclose failure must not promote\n%%EOF"
    session = await service.begin(actor, target, BlobUploadInput("report.pdf", len(data)), "close-failure")
    assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
    with pytest.raises(RuntimeError, match="^BLOB_UPLOAD_INCOMPLETE$") as failure:
        await service.complete(actor, session.upload_id)
    assert "CANARY" not in str(failure.value)
    async with aiosqlite.connect(path) as db:
        async with db.execute("""SELECT u.state,c.status FROM teaql_blob_upload u
            JOIN teaql_blob_catalog c ON c.key=u.key WHERE u.upload_id=?""",
                              (session.upload_id,)) as cursor:
            assert await cursor.fetchone() == ("written", "initiated")
    store.fail_close = False
    assert await service.complete(actor, session.upload_id) == session.reference


@pytest.mark.asyncio
@pytest.mark.parametrize("case", ["fixture-pdf", "real-jpg", "real-jpeg", "real-png", "real-pdf",
                                  "real-docx", "real-xlsx", "real-doc", "real-xls",
                                  "real-mp3", "real-mp4"])
async def test_upload_real_sqlite_s3_create_only_confirm_and_replay(tmp_path, case):
    if os.environ.get("TEAQL_S3_TEST_BLOB_UPLOAD") != "1":
        pytest.skip("explicit S3-compatible Blob upload gate required")
    if case in ("real-jpg", "real-jpeg", "real-png") and os.environ.get("TEAQL_S3_TEST_REAL_IMAGE") != "1":
        pytest.skip("explicit production image S3 gate required")
    if case == "real-pdf" and os.environ.get("TEAQL_S3_TEST_REAL_PDF") != "1":
        pytest.skip("explicit production PDF S3 gate required")
    if case in ("real-docx", "real-xlsx") and os.environ.get("TEAQL_S3_TEST_REAL_OFFICE") != "1":
        pytest.skip("explicit production Office S3 gate required")
    if case in ("real-doc", "real-xls") and os.environ.get("TEAQL_S3_TEST_REAL_LEGACY_OFFICE") != "1":
        pytest.skip("explicit production legacy Office S3 gate required")
    if case in ("real-mp3", "real-mp4") and os.environ.get("TEAQL_S3_TEST_REAL_MEDIA") != "1":
        pytest.skip("explicit production media S3 gate required")
    endpoint = os.environ.get("TEAQL_S3_TEST_ENDPOINT")
    bucket = os.environ.get("TEAQL_S3_TEST_BUCKET")
    access = os.environ.get("TEAQL_S3_TEST_ACCESS_KEY")
    secret = os.environ.get("TEAQL_S3_TEST_SECRET_KEY")
    assert all((endpoint, bucket, access, secret)), "explicit S3 test configuration required"
    from teaql.provider.s3 import S3ObjectStorage

    provider = S3ObjectStorage.connect(bucket, endpoint_url=endpoint,
                                       aws_access_key_id=access, aws_secret_access_key=secret)
    path = str(tmp_path / "s3-upload.db")
    descriptor = EntityDescriptor("Asset").table_name("asset")
    descriptor.property(PropertyDescriptor("id", DataType.I64).is_id())
    descriptor.property(PropertyDescriptor("version", DataType.I64).is_version())
    descriptor.property(PropertyDescriptor("name", DataType.Text))
    descriptor.property(PropertyDescriptor("document", DataType.Text))
    schema = SimpleSchemaProvider()
    schema.register_entity(descriptor)
    actor = RuntimeModule.new().entity(descriptor).into_context()
    actor.insert_resource("dataService", SqlDataServiceExecutor(
        SqliteDialect(), SqliteTransport(path), schema))
    actor.with_object_storage(provider).with_object_storage_observer(lambda event: None)
    clock = [time.time_ns() // 1_000_000]
    catalog = SqliteBlobCatalog(SqliteTransport(path), now_ms=lambda: clock[0])
    if case in ("real-jpg", "real-jpeg", "real-png"):
        from teaql_blob_reference.content_pillow import PillowBlobContentVerifier
        extension = case.removeprefix("real-")
        authority, parser = UploadAuthority((extension,)), PillowBlobContentVerifier()
        fixture = "png" if extension == "png" else "jpg"
        data = (Path(__file__).parents[2] / "fixtures" / ("blob-sample." + fixture)).read_bytes()
        file_name = "sample." + extension
        content_type = "image/png" if extension == "png" else "image/jpeg"
    elif case == "real-pdf":
        from teaql_blob_reference.content_pypdf import PypdfBlobContentVerifier
        authority, parser = UploadAuthority(("pdf",)), PypdfBlobContentVerifier()
        data = (Path(__file__).parents[2] / "fixtures" / "blob-sample.pdf").read_bytes()
        file_name, content_type = "sample.pdf", "application/pdf"
    elif case in ("real-docx", "real-xlsx"):
        from teaql_blob_reference.content_office import OfficeBlobContentVerifier
        extension = case.removeprefix("real-")
        authority, parser = UploadAuthority((extension,)), OfficeBlobContentVerifier()
        output = BytesIO()
        if extension == "docx":
            from docx import Document
            document = Document()
            document.add_paragraph("TeaQL S3 Office upload")
            document.save(output)
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.active["A1"] = "TeaQL S3 Office upload"
            workbook.save(output)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        data, file_name = output.getvalue(), "sample." + extension
    elif case in ("real-doc", "real-xls"):
        from teaql_blob_reference.content_legacy_office import LegacyOfficeBlobContentVerifier
        soffice, bwrap, prlimit = (which(name) for name in ("soffice", "bwrap", "prlimit"))
        if not all((soffice, bwrap, prlimit)):
            pytest.skip("host LibreOffice, bubblewrap and prlimit required")
        extension = case.removeprefix("real-")
        scratch = tmp_path / "legacy-scratch"
        scratch.mkdir(mode=0o700)
        scratch.chmod(0o700)
        authority = UploadAuthority((extension,))
        parser = LegacyOfficeBlobContentVerifier(soffice, bwrap, prlimit, scratch,
                                                 maximum_input_bytes=100_000,
                                                 maximum_converted_bytes=100_000)
        data = (Path(__file__).parents[2] / "fixtures" / "reference-modules" /
                ("sample." + extension)).read_bytes()
        file_name = "sample." + extension
        content_type = "application/msword" if extension == "doc" else "application/vnd.ms-excel"
    elif case in ("real-mp3", "real-mp4"):
        from teaql_blob_reference.content_ffmpeg import FfmpegBlobContentVerifier
        ffprobe, ffmpeg = which("ffprobe"), which("ffmpeg")
        if not ffprobe or not ffmpeg:
            pytest.skip("host FFmpeg command-line tools required")
        extension = case.removeprefix("real-")
        authority = UploadAuthority((extension,))
        parser = FfmpegBlobContentVerifier(ffprobe, ffmpeg)
        data = (Path(__file__).parents[2] / "fixtures" / "reference-modules" /
                ("sample." + extension)).read_bytes()
        file_name = "sample." + extension
        content_type = "audio/mpeg" if extension == "mp3" else "video/mp4"
    else:
        authority, parser = UploadAuthority(), FixturePDFVerifier()
        data = b"%PDF-1.7\nreal S3 Blob upload\n%%EOF"
        file_name, content_type = "report.pdf", "application/pdf"

    class CountingVerifier:
        def __init__(self, inner):
            self.inner = inner
            self.calls = 0

        def verify(self, stream, extension, maximum_bytes):
            self.calls += 1
            return self.inner.verify(stream, extension, maximum_bytes)

    verifier = CountingVerifier(parser)
    service = SqliteBlobUploadService(catalog, authority, verifier, 100_000,
                                      now_ms=lambda: clock[0])
    object_key = None
    try:
        async with aiosqlite.connect(path) as db:
            await db.execute("CREATE TABLE asset (id INTEGER PRIMARY KEY, version INTEGER, name TEXT, document TEXT)")
            await db.commit()
        await catalog.ensure_schema()
        session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                      BlobUploadInput(file_name, len(data), content_type),
                                      uuid4().hex)
        async with aiosqlite.connect(path) as db:
            async with db.execute("SELECT object_key FROM teaql_blob_catalog WHERE key=?",
                                  (session.reference.key,)) as cursor:
                object_key = ObjectKey((await cursor.fetchone())[0])
        assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
        authority.allowed = False
        with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
            await service.complete(actor, session.upload_id)
        authority.allowed = True
        assert await service.complete(actor, session.upload_id) == session.reference
        assert await service.complete(actor, session.upload_id) == session.reference
        assert verifier.calls == 1
        assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
        store = actor.object_storage("verify completed S3 Blob upload")
        assert store.get_bytes(object_key, 100_000) == data
        head = store.head(object_key)
        async with aiosqlite.connect(path) as db:
            async with db.execute(
                "SELECT status, content_type, size_bytes, content_version "
                "FROM teaql_blob_catalog WHERE key=?", (session.reference.key,)
            ) as cursor:
                status, mime, size, version = await cursor.fetchone()
        assert (status, mime, size, version) == (
            "uploaded_unbound", content_type, len(data), head.etag)

        class LifecycleBindingAuthority(BindingAuthority):
            async def authorize(self, _context, target, previous, next_ref, action):
                return (target == BlobBindingTarget("Asset", "1", "document")
                        and ((action == "Create" and previous is None
                              and next_ref == session.reference)
                             or (action == "Update" and previous == session.reference
                                 and next_ref is None)))

        binding_authority = LifecycleBindingAuthority()

        async def bind(graph):
            command = (InsertCommand.new("Asset").value("id", 1)
                       .value("version", 1).value("name", "document")
                       .value("document", session.reference.key))
            await graph.transaction.mutate(graph.context,
                graph.request(command, graph.scope(EntityKey("Asset", 1))))
            assert await catalog.bind_in_graph(graph, graph.context,
                BlobBindingChange(BlobBindingTarget("Asset", "1", "document"),
                                  session.reference, 0, 1, "Create"),
                binding_authority) == 1

        await actor.execute_graph_save(bind, comment="bind S3 Blob in business graph")
        async with aiosqlite.connect(path) as db:
            async with db.execute("SELECT document,version FROM asset WHERE id=1") as cursor:
                assert await cursor.fetchone() == (session.reference.key, 1)
            async with db.execute("SELECT status FROM teaql_blob_catalog WHERE key=?",
                                  (session.reference.key,)) as cursor:
                assert (await cursor.fetchone())[0] == "available"

        class ReadAuthorization:
            allowed = True
            async def authorize(self, _context, _slot):
                return self.allowed

        class Authorizer:
            version = 1
            async def authorize(self, _context, _result, _options):
                slot = AuthorizedBlobSlot(("document",), "Asset", "1", "document",
                                          "tenant-a", self.version, session.reference)
                return AuthorizedResponseGraph({"document": session.reference}, (slot,))

        read_permission, authorizer = ReadAuthorization(), Authorizer()
        actor.insert_resource("response_graph_authorizer", authorizer)
        actor.insert_resource("blob_response_contributor",
                              SqliteBlobResponseContributor(catalog, read_permission, 3))
        options = ViewOptions("ordinary", "read_only",
                              datetime.now(timezone.utc) + timedelta(seconds=60))
        view = await ViewPreparationService().prepare(actor, None, options)
        link = view.presentation["document"]
        assert link.startswith("http") and view.cache_control == "private, no-store"
        with urlopen(Request(link, headers={"Range": "bytes=0-7"}), timeout=5) as response:
            assert response.status == 206 and response.read() == data[:8]
            assert response.headers["Content-Type"] == content_type
            assert file_name in response.headers["Content-Disposition"]
        read_permission.allowed = False
        with pytest.raises(ViewPreparationError) as denied:
            await ViewPreparationService().prepare(actor, None, options)
        assert denied.value.code == "BLOB_READ_DENIED"
        read_permission.allowed = True
        authorizer.version = 0
        with pytest.raises(ViewPreparationError) as stale:
            await ViewPreparationService().prepare(actor, None,
                ViewOptions("ordinary", "committed",
                            datetime.now(timezone.utc) + timedelta(seconds=60)))
        assert (stale.value.code, stale.value.business_state) == ("BLOB_READ_DENIED", "committed")
        if os.environ.get("TEAQL_S3_TEST_END_TO_END_GC") == "1":
            async def unbind(graph):
                command = (UpdateCommand.new("Asset", 1).expected_version(1)
                           .value("document", None))
                await graph.transaction.mutate(graph.context,
                    graph.request(command, graph.scope(EntityKey("Asset", 1))))
                assert await catalog.bind_in_graph(graph, graph.context,
                    BlobBindingChange(BlobBindingTarget("Asset", "1", "document"),
                                      None, 1, 2, "Update"), binding_authority) == 2

            await actor.execute_graph_save(unbind, comment="unbind S3 Blob in business graph")
            async with aiosqlite.connect(path) as db:
                async with db.execute("SELECT document,version FROM asset WHERE id=1") as cursor:
                    assert await cursor.fetchone() == (None, 2)
                async with db.execute("SELECT key,version FROM teaql_blob_binding WHERE "
                                      "scope='tenant-a' AND entity='Asset' AND entity_id='1' "
                                      "AND field='document'") as cursor:
                    assert await cursor.fetchone() == (None, 2)
                async with db.execute(
                    "SELECT status,last_unbound_at_ms,effective_unref_seconds,"
                    "retention_floor_until_ms FROM teaql_blob_catalog WHERE key=?",
                    (session.reference.key,)
                ) as cursor:
                    status, last_unbound, effective, floor = await cursor.fetchone()
            assert status == "available" and last_unbound == clock[0]
            assert effective == 604800 and floor >= clock[0] + 3600 * 1000
            authorizer.version = 1
            with pytest.raises(ViewPreparationError) as unbound_view:
                await ViewPreparationService().prepare(actor, None, options)
            assert unbound_view.value.code == "BLOB_READ_DENIED"

            class MaintenanceAuthority:
                def scope(self, _context):
                    return "tenant-a"

                async def can_maintain(self, _context, scope):
                    return scope == "tenant-a"

            maintenance = SqliteBlobMaintenanceService(
                catalog, MaintenanceAuthority(), 10, now_ms=lambda: clock[0])
            deadline = max(last_unbound + effective * 1000, floor)
            clock[0] = deadline - 1
            assert (await maintenance.run_batch(actor)).outcomes == ()
            assert store.head(object_key) is not None
            clock[0] = deadline
            collected = await maintenance.run_batch(actor)
            assert [(item.reference, item.status) for item in collected.outcomes] == [
                (session.reference, "deleted")]
            assert store.head(object_key) is None
            async with aiosqlite.connect(path) as db:
                async with db.execute("SELECT status FROM teaql_blob_catalog WHERE key=?",
                                      (session.reference.key,)) as cursor:
                    assert (await cursor.fetchone())[0] == "deleted"
            print("PASS Python " + case +
                  " S3 upload/confirm/graph bind/authorized view/unbind/retention GC")
    finally:
        try:
            if object_key is not None:
                actor.object_storage("clean S3 Blob upload fixture").delete(object_key)
        finally:
            provider.close()


@pytest.mark.asyncio
async def test_upload_rechecks_current_permission_on_physical_sql_lease(tmp_path):
    path = str(tmp_path / "lease-authority.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    async with aiosqlite.connect(path) as db:
        await db.execute("CREATE TABLE upload_permission (id INTEGER PRIMARY KEY, allowed INTEGER)")
        await db.execute("INSERT INTO upload_permission VALUES (1,0)")
        await db.commit()

    class LeaseAuthority(UploadAuthority):
        def __init__(self):
            super().__init__()
            self.seen = []

        async def current_on_lease(self, actor, lease, target, action, _input):
            self.seen.append(action)
            rows = await lease.fetch_all_sql(CompiledQuery(
                "SELECT allowed FROM upload_permission WHERE id=1", []))
            assert len(rows) == 1
            return BlobUploadAuthorizationState(self.scope(actor), self.field(actor, target),
                                                rows[0]["allowed"] == 1)

    async def set_allowed(value):
        async with aiosqlite.connect(path) as db:
            await db.execute("UPDATE upload_permission SET allowed=? WHERE id=1", (value,))
            await db.commit()

    authority = LeaseAuthority()
    verifier = FixturePDFVerifier()
    service = SqliteBlobUploadService(catalog, authority, verifier, 10_000)
    actor = RuntimeModule.new().into_context().with_object_storage(
        FileObjectStorage(tmp_path / "private", os.urandom(32)))
    actor.with_object_storage_observer(lambda event: None)
    target = BlobUploadTarget("Asset", "1", "document")
    data = b"%PDF-1.7\nlease permission\n%%EOF"
    input = BlobUploadInput("report.pdf", len(data), "application/pdf")

    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.begin(actor, target, input, "lease-denied")
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT COUNT(*) FROM teaql_blob_upload") as cursor:
            assert (await cursor.fetchone())[0] == 0

    await set_allowed(1)
    session = await service.begin(actor, target, input, "lease-allowed")
    await set_allowed(0)
    source = BytesIO(data)
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.write(actor, session.upload_id, source)
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT state FROM teaql_blob_upload WHERE upload_id=?",
                              (session.upload_id,)) as cursor:
            assert (await cursor.fetchone())[0] == "initiated"
    await set_allowed(1)
    assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
    await set_allowed(0)
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.complete(actor, session.upload_id)
    assert verifier.calls == 0
    await set_allowed(1)
    assert await service.complete(actor, session.upload_id) == session.reference
    assert verifier.calls == 1
    await set_allowed(0)
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.complete(actor, session.upload_id)
    assert verifier.calls == 1
    await set_allowed(1)
    assert await service.complete(actor, session.upload_id) == session.reference
    assert verifier.calls == 1
    assert {"Begin", "Write", "Confirm"}.issubset(authority.seen)


@pytest.mark.asyncio
async def test_expired_replay_key_cannot_recover_upload_session(tmp_path):
    catalog = SqliteBlobCatalog(SqliteTransport(str(tmp_path / "expired.db")))
    await catalog.ensure_schema()
    clock = [1_700_000_000_000]
    service = SqliteBlobUploadService(catalog, UploadAuthority(), FixturePDFVerifier(),
                                      10_000, lease_seconds=1, now_ms=lambda: clock[0])
    context = RuntimeModule.new().into_context()
    target = BlobUploadTarget("Asset", "1", "document")
    input = BlobUploadInput("report.pdf", 42)
    first = await service.begin(context, target, input, "expired-request")
    assert (await service.begin(context, target, input, "expired-request")).upload_id == first.upload_id
    clock[0] = int(first.expires_at.timestamp() * 1000)
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_EXPIRED"):
        await service.begin(context, target, input, "expired-request")
    second = await service.begin(context, target, input, "fresh-request")
    assert second.upload_id != first.upload_id
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_EXPIRED"):
        await service.write(context, first.upload_id, BytesIO(b"x" * 42))


@pytest.mark.asyncio
async def test_upload_write_holds_sql_owner_across_file_io(tmp_path):
    path = str(tmp_path / "locked.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    entered, release = Event(), Event()
    file_store = FileObjectStorage(tmp_path / "private", os.urandom(32))

    class BlockingStore:
        puts = 0

        def __getattr__(self, name):
            return getattr(file_store, name)

        def put(self, invocation, key, content, options):
            self.puts += 1
            result = file_store.put(invocation, key, content, options)
            entered.set()
            if not release.wait(5):
                raise AssertionError("test did not release File write")
            return result

    store = BlockingStore()
    actor = RuntimeModule.new().into_context().with_object_storage(store)
    actor.with_object_storage_observer(lambda event: None)
    service = SqliteBlobUploadService(catalog, UploadAuthority(), FixturePDFVerifier(), 10_000)
    data = b"%PDF-1.7\nlocked upload\n%%EOF"
    session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                  BlobUploadInput("report.pdf", len(data)), "lock-request")
    first = asyncio.create_task(service.write(actor, session.upload_id, BytesIO(data)))
    second = None
    try:
        assert await asyncio.to_thread(entered.wait, 5)
        second = asyncio.create_task(service.write(actor, session.upload_id, BytesIO(data)))
        await asyncio.sleep(0.05)
        assert not second.done(), "second writer bypassed the physical SQL owner"
    finally:
        release.set()
    assert await first == len(data)
    assert await second == len(data)
    assert store.puts == 1
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT state FROM teaql_blob_upload WHERE upload_id=?",
                              (session.upload_id,)) as cursor:
            assert (await cursor.fetchone())[0] == "written"


@pytest.mark.asyncio
async def test_upload_confirmation_holds_sql_owner_through_parser_io(tmp_path):
    path = str(tmp_path / "confirm-locked.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    entered, release = Event(), Event()

    class BlockingVerifier(FixturePDFVerifier):
        def verify(self, stream, extension, maximum_bytes):
            result = super().verify(stream, extension, maximum_bytes)
            entered.set()
            if not release.wait(5):
                raise AssertionError("test did not release parser")
            return result

    verifier = BlockingVerifier()
    actor = RuntimeModule.new().into_context().with_object_storage(
        FileObjectStorage(tmp_path / "private", os.urandom(32)))
    actor.with_object_storage_observer(lambda event: None)
    service = SqliteBlobUploadService(catalog, UploadAuthority(), verifier, 10_000)
    data = b"%PDF-1.7\nlocked confirmation\n%%EOF"
    session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                  BlobUploadInput("report.pdf", len(data)), "confirm-request")
    await service.write(actor, session.upload_id, BytesIO(data))
    first = asyncio.create_task(service.complete(actor, session.upload_id))
    second = None
    try:
        assert await asyncio.to_thread(entered.wait, 5)
        second = asyncio.create_task(service.complete(actor, session.upload_id))
        await asyncio.sleep(0.05)
        assert not second.done(), "second confirmation bypassed the physical SQL owner"
    finally:
        release.set()
    assert await first == session.reference
    assert await second == session.reference
    assert await service.complete(actor, session.upload_id) == session.reference
    assert verifier.calls == 1


@pytest.mark.asyncio
async def test_confirm_serializes_permission_revocation_during_parser_io(tmp_path):
    path = str(tmp_path / "confirm-revocation.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    async with aiosqlite.connect(path) as db:
        await db.execute("CREATE TABLE upload_permission (id INTEGER PRIMARY KEY, allowed INTEGER)")
        await db.execute("INSERT INTO upload_permission VALUES (1,1)")
        await db.commit()
    entered, release = Event(), Event()

    class LeaseAuthority(UploadAuthority):
        async def current_on_lease(self, actor, lease, target, action, _input):
            rows = await lease.fetch_all_sql(CompiledQuery(
                "SELECT allowed FROM upload_permission WHERE id=1", []))
            return BlobUploadAuthorizationState(self.scope(actor), self.field(actor, target),
                                                len(rows) == 1 and rows[0]["allowed"] == 1)

    class BlockingVerifier(FixturePDFVerifier):
        def verify(self, stream, extension, maximum_bytes):
            result = super().verify(stream, extension, maximum_bytes)
            entered.set()
            if not release.wait(5):
                raise AssertionError("test did not release parser")
            return result

    actor = RuntimeModule.new().into_context().with_object_storage(
        FileObjectStorage(tmp_path / "private", os.urandom(32)))
    actor.with_object_storage_observer(lambda event: None)
    verifier = BlockingVerifier()
    service = SqliteBlobUploadService(catalog, LeaseAuthority(), verifier, 10_000)
    data = b"%PDF-1.7\nconcurrent revocation\n%%EOF"
    session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                  BlobUploadInput("report.pdf", len(data)), "revoke-request")
    assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)

    async def revoke():
        async with aiosqlite.connect(path) as db:
            await db.execute("UPDATE upload_permission SET allowed=0 WHERE id=1")
            await db.commit()

    confirmation = asyncio.create_task(service.complete(actor, session.upload_id))
    revocation = None
    try:
        assert await asyncio.to_thread(entered.wait, 5)
        revocation = asyncio.create_task(revoke())
        await asyncio.sleep(0.05)
        assert not revocation.done(), "revocation committed during the parser's SQL owner"
    finally:
        release.set()
    assert await confirmation == session.reference
    assert revocation is not None
    await revocation
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.complete(actor, session.upload_id)
    assert verifier.calls == 1
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT state FROM teaql_blob_upload WHERE upload_id=?",
                              (session.upload_id,)) as cursor:
            assert (await cursor.fetchone())[0] == "completed"


@pytest.mark.asyncio
async def test_upload_write_commit_unknown_replays_only_same_bytes_and_generation(tmp_path):
    path = str(tmp_path / "commit-unknown.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    file_store = FileObjectStorage(tmp_path / "private", os.urandom(32))

    class CountingStore:
        puts = 0

        def __getattr__(self, name):
            return getattr(file_store, name)

        def put(self, invocation, key, content, options):
            self.puts += 1
            return file_store.put(invocation, key, content, options)

    store = CountingStore()
    actor = RuntimeModule.new().into_context().with_object_storage(store)
    actor.with_object_storage_observer(lambda event: None)
    authority, verifier = UploadAuthority(), FixturePDFVerifier()
    service = SqliteBlobUploadService(catalog, authority, verifier, 10_000)
    data = b"%PDF-1.7\nSQL commit succeeded, reply lost\n%%EOF"
    session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                  BlobUploadInput("report.pdf", len(data)), "unknown-write")
    begin_sql = catalog.transport.begin_sql

    class CommitThenFail:
        def __init__(self, tx):
            self.tx = tx

        def __getattr__(self, name):
            return getattr(self.tx, name)

        async def commit_sql(self):
            await self.tx.commit_sql()
            raise OSError("PRIVATE-COMMIT-CANARY")

    async def uncertain_begin():
        return CommitThenFail(await begin_sql())

    catalog.transport.begin_sql = uncertain_begin
    try:
        with pytest.raises(RuntimeError, match="^BLOB_CATALOG_FAILED$"):
            await service.write(actor, session.upload_id, BytesIO(data))
    finally:
        catalog.transport.begin_sql = begin_sql
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT u.state,c.status,c.object_key,u.written_etag "
                              "FROM teaql_blob_upload u JOIN teaql_blob_catalog c ON c.key=u.key "
                              "WHERE u.upload_id=?", (session.upload_id,)) as cursor:
            state, catalog_state, object_key, etag = await cursor.fetchone()
    assert state == "written" and catalog_state == "initiated" and etag
    assert store.puts == 1
    authority.allowed = False
    with pytest.raises(RuntimeError, match="^BLOB_UPLOAD_DENIED$"):
        await service.write(actor, session.upload_id, BytesIO(data))
    authority.allowed = True
    with pytest.raises(RuntimeError, match="^BLOB_IDEMPOTENCY_CONFLICT$"):
        await service.write(actor, session.upload_id, BytesIO(b"x" * len(data)))
    class FailingSource(BytesIO):
        def read(self, *_):
            raise OSError("PRIVATE-SOURCE-CANARY")

    with pytest.raises(RuntimeError, match="^BLOB_STORE_FAILED$") as failed:
        await service.write(actor, session.upload_id, FailingSource(data))
    assert "CANARY" not in str(failed.value)
    assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
    assert store.puts == 1
    assert actor.object_storage("check unchanged version").head(ObjectKey(object_key)).etag == etag
    catalog.transport.begin_sql = uncertain_begin
    try:
        with pytest.raises(RuntimeError, match="^BLOB_CATALOG_FAILED$"):
            await service.complete(actor, session.upload_id)
    finally:
        catalog.transport.begin_sql = begin_sql
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT state FROM teaql_blob_upload WHERE upload_id=?",
                              (session.upload_id,)) as cursor:
            assert (await cursor.fetchone())[0] == "completed"
    authority.allowed = False
    with pytest.raises(RuntimeError, match="^BLOB_UPLOAD_DENIED$"):
        await service.complete(actor, session.upload_id)
    authority.allowed = True
    assert await service.complete(actor, session.upload_id) == session.reference
    assert verifier.calls == 1
    assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
    assert verifier.calls == 1 and store.puts == 1
    actor.object_storage("replace generation").put_bytes(ObjectKey(object_key), data,
        PutOptions(len(data), WriteMode.REPLACE, "application/pdf"))
    with pytest.raises(RuntimeError, match="^BLOB_IDEMPOTENCY_CONFLICT$"):
        await service.write(actor, session.upload_id, BytesIO(data))


@pytest.mark.asyncio
async def test_upload_retry_recovers_only_identical_file_after_journal_rollback(tmp_path):
    path = str(tmp_path / "recover.db")
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    await catalog.ensure_schema()
    actor = RuntimeModule.new().into_context().with_object_storage(
        FileObjectStorage(tmp_path / "private", os.urandom(32)))
    actor.with_object_storage_observer(lambda event: None)
    verifier = FixturePDFVerifier()
    service = SqliteBlobUploadService(catalog, UploadAuthority(), verifier, 10_000)
    data = b"%PDF-1.7\nrecover committed File content\n%%EOF"
    session = await service.begin(actor, BlobUploadTarget("Asset", "1", "document"),
                                  BlobUploadInput("report.pdf", len(data)), "recovery-request")
    async with aiosqlite.connect(path) as db:
        await db.execute("""CREATE TRIGGER reject_upload_write BEFORE UPDATE ON teaql_blob_upload
            WHEN NEW.state='written' BEGIN SELECT RAISE(ABORT, 'injected'); END""")
        await db.commit()
    with pytest.raises(RuntimeError, match="BLOB_CATALOG_FAILED"):
        await service.write(actor, session.upload_id, BytesIO(data))
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT c.object_key,u.state FROM teaql_blob_upload u "
                              "JOIN teaql_blob_catalog c ON c.key=u.key WHERE u.upload_id=?",
                              (session.upload_id,)) as cursor:
            object_key, state = await cursor.fetchone()
        assert state == "initiated"
        await db.execute("DROP TRIGGER reject_upload_write")
        await db.commit()
    stored = actor.object_storage("check recovery candidate").head(ObjectKey(object_key))
    assert stored is not None and stored.size == len(data)
    with pytest.raises(RuntimeError, match="BLOB_STORE_FAILED"):
        await service.write(actor, session.upload_id, BytesIO(b"x" * len(data)))
    assert await service.write(actor, session.upload_id, BytesIO(data)) == len(data)
    assert actor.object_storage("check recovered generation").head(
        ObjectKey(object_key)).etag == stored.etag
    assert await service.complete(actor, session.upload_id) == session.reference
    assert verifier.calls == 1


@pytest.mark.asyncio
async def test_upload_real_sqlite_file_verifier_and_native_bind(tmp_path):
    path = str(tmp_path / "upload.db")
    descriptor = EntityDescriptor("Asset").table_name("asset")
    descriptor.property(PropertyDescriptor("id", DataType.I64).is_id())
    descriptor.property(PropertyDescriptor("version", DataType.I64).is_version())
    descriptor.property(PropertyDescriptor("name", DataType.Text))
    schema = SimpleSchemaProvider()
    schema.register_entity(descriptor)
    context = RuntimeModule.new().entity(descriptor).into_context()
    context.insert_resource("dataService", SqlDataServiceExecutor(
        SqliteDialect(), SqliteTransport(path), schema))
    provider = FileObjectStorage(tmp_path / "private", os.urandom(32))
    context.with_object_storage(provider).with_object_storage_observer(lambda event: None)
    catalog = SqliteBlobCatalog(SqliteTransport(path))
    async with aiosqlite.connect(path) as db:
        await db.execute("CREATE TABLE asset (id INTEGER PRIMARY KEY, version INTEGER, name TEXT)")
        await db.commit()
    await catalog.ensure_schema()
    await catalog.validate_schema()
    authority = UploadAuthority()
    verifier = FixturePDFVerifier()
    service = SqliteBlobUploadService(catalog, authority, verifier, 10_000)
    target = BlobUploadTarget("Asset", "1", "document")
    data = b"%PDF-1.7\nreal file bytes\n%%EOF"
    input = BlobUploadInput("path/report.pdf", len(data), "application/pdf")

    session = await service.begin(context, target, input, "request-1")
    replay = await service.begin(context, target, input, "request-1")
    assert replay.upload_id == session.upload_id
    with pytest.raises(RuntimeError, match="BLOB_IDEMPOTENCY_CONFLICT"):
        await service.begin(context, target, BlobUploadInput("other.pdf", len(data)), "request-1")
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_INCOMPLETE"):
        await service.complete(context, session.upload_id)
    authority.current_scope = "tenant-b"
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.write(context, session.upload_id, BytesIO(data))
    authority.current_scope = "tenant-a"
    assert await service.write(context, session.upload_id, BytesIO(data)) == len(data)
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT status FROM teaql_blob_catalog WHERE key=?",
                              (session.reference.key,)) as cursor:
            assert (await cursor.fetchone())[0] == "initiated"
    authority.allowed = False
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.complete(context, session.upload_id)
    authority.allowed = True
    reference = await service.complete(context, session.upload_id)
    assert reference == session.reference and verifier.calls == 1
    assert await service.complete(context, session.upload_id) == reference
    assert verifier.calls == 1
    authority.allowed = False
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_DENIED"):
        await service.complete(context, session.upload_id)
    authority.allowed = True

    async def bind(graph):
        command = (InsertCommand.new("Asset").value("id", 1)
                   .value("version", 1).value("name", "document"))
        await graph.transaction.mutate(graph.context,
            graph.request(command, graph.scope(EntityKey("Asset", 1))))
        await catalog.bind_in_graph(graph, graph.context,
            BlobBindingChange(BlobBindingTarget("Asset", "1", "document"),
                              reference, 0, 1, "Create"), BindingAuthority())

    await context.execute_graph_save(bind, comment="bind verified upload")
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT status FROM teaql_blob_catalog WHERE key=?",
                              (reference.key,)) as cursor:
            assert (await cursor.fetchone())[0] == "available"
    assert await service.complete(context, session.upload_id) == reference
    assert verifier.calls == 1
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT object_key FROM teaql_blob_catalog WHERE key=?",
                              (reference.key,)) as cursor:
            completed_key = (await cursor.fetchone())[0]
    context.object_storage("replace completed generation").put_bytes(
        ObjectKey(completed_key), data, PutOptions(len(data), WriteMode.REPLACE, "application/pdf"))
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_INCOMPLETE"):
        await service.complete(context, session.upload_id)
    assert verifier.calls == 1

    second = await service.begin(context, target, input, "request-2")
    await service.write(context, second.upload_id, BytesIO(data))
    async with aiosqlite.connect(path) as db:
        async with db.execute("SELECT object_key FROM teaql_blob_catalog WHERE key=?",
                              (second.reference.key,)) as cursor:
            object_key = (await cursor.fetchone())[0]
    context.object_storage("replace generation").put_bytes(
        ObjectKey(object_key), data, PutOptions(len(data), WriteMode.REPLACE, "application/pdf"))
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_INCOMPLETE"):
        await service.complete(context, second.upload_id)

    third = await service.begin(context, target, input, "request-3")
    await service.write(context, third.upload_id, BytesIO(data))
    early = SqliteBlobUploadService(catalog, authority, FixturePDFVerifier(early=True), 10_000)
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_INCOMPLETE"):
        await early.complete(context, third.upload_id)

    fourth = await service.begin(context, target, input, "request-4")
    with pytest.raises(RuntimeError, match="BLOB_STORE_FAILED"):
        await service.write(context, fourth.upload_id, BytesIO(data + b"x"))
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_INCOMPLETE"):
        await service.complete(context, fourth.upload_id)
    fifth = await service.begin(context, target, input, "request-5")
    with pytest.raises(RuntimeError, match="BLOB_STORE_FAILED"):
        await service.write(context, fifth.upload_id, BytesIO(data[:-1]))
    with pytest.raises(RuntimeError, match="BLOB_UPLOAD_INCOMPLETE"):
        await service.complete(context, fifth.upload_id)
