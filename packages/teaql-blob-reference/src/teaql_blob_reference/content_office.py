"""Optional bounded DOCX/XLSX parser for Blob upload verification."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from threading import BoundedSemaphore
from xml.etree import ElementTree
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile

from .upload_service import BlobVerifiedContent


_MIME = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
_REQUIRED = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml"}
_MAIN_TYPE = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
}


class OfficeBlobContentVerifier:
    """Host-selected Office parser; construction does no parser import or storage I/O."""

    supported_extensions = frozenset(_MIME)

    def __init__(self, maximum_input_bytes: int = 16 * 1024 * 1024,
                 maximum_decoded_bytes: int = 32 * 1024 * 1024,
                 maximum_parts: int = 1000, maximum_elements: int = 1_000_000,
                 concurrent_verifications: int = 2):
        if (type(maximum_input_bytes) is not int or not 1 <= maximum_input_bytes <= 100 * 1024 * 1024
                or type(maximum_decoded_bytes) is not int
                or not 1 <= maximum_decoded_bytes <= 100 * 1024 * 1024
                or type(maximum_parts) is not int or not 1 <= maximum_parts <= 10000
                or type(maximum_elements) is not int or not 1 <= maximum_elements <= 2_000_000
                or type(concurrent_verifications) is not int
                or not 1 <= concurrent_verifications <= 32):
            raise RuntimeError("BLOB_POLICY_INVALID")
        self.maximum_input_bytes = maximum_input_bytes
        self.maximum_decoded_bytes = maximum_decoded_bytes
        self.maximum_parts = maximum_parts
        self.maximum_elements = maximum_elements
        self._permits = BoundedSemaphore(concurrent_verifications)

    def validate_ready(self) -> None:
        try:
            from docx import Document  # noqa: F401
            from openpyxl import load_workbook  # noqa: F401
        except ImportError:
            raise RuntimeError("BLOB_CAPABILITY_MISSING") from None

    def _check_package(self, data: bytes, extension: str) -> None:
        try:
            with ZipFile(BytesIO(data)) as archive:
                parts = archive.infolist()
                if not 1 <= len(parts) <= self.maximum_parts:
                    raise ValueError("Office part count")
                names = set()
                for part in parts:
                    name = part.filename
                    if (name in names or name.startswith("/") or "\\" in name
                            or any(piece in ("", ".", "..") for piece in name.split("/"))
                            or part.flag_bits & 1 or part.compress_type not in (ZIP_STORED, ZIP_DEFLATED)):
                        raise ValueError("invalid Office member")
                    names.add(name)
                if not {"[Content_Types].xml", "_rels/.rels", _REQUIRED[extension]} <= names:
                    raise ValueError("missing Office document parts")
                decoded = 0
                for part in parts:
                    if part.is_dir():
                        continue
                    if part.file_size > self.maximum_decoded_bytes - decoded:
                        raise ValueError("Office decoded size")
                    with archive.open(part) as member:
                        while True:
                            block = member.read(65536)
                            if not block:
                                break
                            decoded += len(block)
                            if decoded > self.maximum_decoded_bytes:
                                raise ValueError("Office decoded size")
                    # Reading through EOF also verifies the member's CRC.
                types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
                main = "/" + _REQUIRED[extension]
                if not any(item.attrib.get("PartName") == main
                           and item.attrib.get("ContentType") == _MAIN_TYPE[extension]
                           for item in types):
                    raise ValueError("wrong Office main content type")
                relationships = ElementTree.fromstring(archive.read("_rels/.rels"))
                if not any(item.attrib.get("Type", "").endswith("/officeDocument")
                           and item.attrib.get("Target") == _REQUIRED[extension]
                           and item.attrib.get("TargetMode", "Internal") == "Internal"
                           for item in relationships):
                    raise ValueError("wrong Office root relationship")
        except (BadZipFile, OSError, EOFError, RuntimeError) as error:
            raise ValueError("invalid Office package") from error

    def _parse_docx(self, data: bytes) -> None:
        from docx import Document

        document = Document(BytesIO(data))
        if document.element.body is None:
            raise ValueError("missing Word body")
        count = 0
        for _ in document.element.body.iter():
            count += 1
            if count > self.maximum_elements:
                raise ValueError("Word element limit")

    def _parse_xlsx(self, data: bytes) -> None:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False,
                                 keep_links=False)
        try:
            if not workbook.worksheets:
                raise ValueError("missing Excel worksheet")
            count = 0
            for sheet in workbook.worksheets:
                sheet.reset_dimensions()  # Do not trust producer-supplied used range.
                for row in sheet.iter_rows():
                    count += len(row)
                    if count > self.maximum_elements:
                        raise ValueError("Excel cell limit")
                    for cell in row:
                        _ = cell.value
        finally:
            workbook.close()

    def verify(self, stream, extension: str, maximum_bytes: int) -> BlobVerifiedContent:
        if extension not in _MIME:
            raise RuntimeError("BLOB_FORMAT_NOT_ALLOWED")
        if type(maximum_bytes) is not int or maximum_bytes < 1:
            raise RuntimeError("BLOB_POLICY_INVALID")
        if not self._permits.acquire(blocking=False):
            raise RuntimeError("BLOB_VERIFICATION_BUSY")
        try:
            maximum = min(maximum_bytes, self.maximum_input_bytes)
            chunks = []
            size = 0
            digest = sha256()
            while True:
                try:
                    block = stream.read(65536)
                except Exception:
                    raise RuntimeError("BLOB_STORE_FAILED") from None
                if not isinstance(block, bytes):
                    raise RuntimeError("BLOB_STORE_FAILED")
                if not block:
                    break
                if len(block) > maximum - size:
                    raise RuntimeError("BLOB_SIZE_EXCEEDED")
                chunks.append(block)
                digest.update(block)
                size += len(block)
            data = b"".join(chunks)
            try:
                self.validate_ready()
                self._check_package(data, extension)
                if extension == "docx":
                    self._parse_docx(data)
                else:
                    self._parse_xlsx(data)
            except RuntimeError as error:
                if str(error) == "BLOB_CAPABILITY_MISSING":
                    raise
                raise RuntimeError("BLOB_FORMAT_NOT_ALLOWED") from None
            except Exception:
                raise RuntimeError("BLOB_FORMAT_NOT_ALLOWED") from None
            return BlobVerifiedContent(size, digest.hexdigest(), _MIME[extension])
        finally:
            self._permits.release()
